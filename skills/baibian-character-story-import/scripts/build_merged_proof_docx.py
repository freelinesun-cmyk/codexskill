#!/usr/bin/env python3
"""从 baibian.characterStories.v1 JSON 生成合并校对 DOCX。"""

from __future__ import annotations

import argparse
import json
import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


RGB_RE = re.compile(r"color:\s*rgb\(\s*(\d{1,3})\s*,\s*(\d{1,3})\s*,\s*(\d{1,3})\s*\)", re.I)


class StoryHtmlParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.paragraphs: list[list[tuple[str, bool, tuple[int, int, int] | None]]] = []
        self.current: list[tuple[str, bool, tuple[int, int, int] | None]] = []
        self.bold_depth = 0
        self.colors: list[tuple[int, int, int] | None] = [None]
        self.in_div = False

    def flush(self, force: bool = False) -> None:
        if self.current or force:
            self.paragraphs.append(self.current)
        self.current = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag == "div":
            if self.in_div or self.current:
                self.flush()
            self.in_div = True
        elif tag == "br":
            self.current.append(("\n", self.bold_depth > 0, self.colors[-1]))
        elif tag == "b":
            self.bold_depth += 1
        elif tag == "span":
            style = dict(attrs).get("style") or ""
            match = RGB_RE.search(style)
            color = tuple(min(255, int(value)) for value in match.groups()) if match else self.colors[-1]
            self.colors.append(color)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag == "div":
            self.flush(force=True)
            self.in_div = False
        elif tag == "b":
            self.bold_depth = max(0, self.bold_depth - 1)
        elif tag == "span" and len(self.colors) > 1:
            self.colors.pop()

    def handle_data(self, data: str) -> None:
        if data:
            self.current.append((data, self.bold_depth > 0, self.colors[-1]))

    def finish(self) -> list[list[tuple[str, bool, tuple[int, int, int] | None]]]:
        if self.current:
            self.flush()
        return self.paragraphs


def set_run_font(run, font_name: str, size: float) -> None:
    run.font.name = font_name
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)


def add_story_paragraph(document: Document, segments, font_name: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.35
    if not segments:
        return
    for text, bold, color in segments:
        pieces = text.split("\n")
        for index, piece in enumerate(pieces):
            if piece:
                run = paragraph.add_run(piece)
                set_run_font(run, font_name, 11.5)
                run.bold = bold
                if color is not None:
                    run.font.color.rgb = RGBColor(*color)
            if index < len(pieces) - 1:
                paragraph.add_run().add_break(WD_BREAK.LINE)


def record_paragraphs(record: dict):
    html = str(record.get("contentHtml") or "")
    if html.strip():
        parser = StoryHtmlParser()
        parser.feed(html)
        parser.close()
        return parser.finish()
    text = str(record.get("contentText") or "")
    return [[(line, False, None)] if line else [] for line in text.splitlines()]


def role_label(roles) -> str:
    values = [str(role).strip() for role in roles or [] if str(role).strip()]
    return values[0] if len(values) == 1 else "全部角色"


def build(input_path: Path, output_path: Path) -> int:
    payload = json.loads(input_path.read_text(encoding="utf-8"))
    if payload.get("schema") != "baibian.characterStories.v1":
        raise ValueError("JSON schema 应为 baibian.characterStories.v1")
    records = payload.get("records")
    if not isinstance(records, list) or not records:
        raise ValueError("records 必须是非空数组")

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    normal = document.styles["Normal"]
    normal.font.name = "Hiragino Sans GB"
    normal.font.size = Pt(11.5)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    heading = document.styles["Heading 1"]
    heading.font.name = "Hiragino Sans GB"
    heading.font.size = Pt(17)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")

    for index, record in enumerate(records, 1):
        if index > 1:
            document.add_page_break()
        title = str(record.get("title") or "未命名故事").strip()
        role = role_label(record.get("roles"))
        document.add_heading(f"{index}  {role} / {title}", level=1)
        for segments in record_paragraphs(record):
            add_story_paragraph(document, segments, "Hiragino Sans GB")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = f"{input_path.stem} 合并校对稿"
    document.save(output_path)
    return len(records)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("json_file", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    source = args.json_file.expanduser().resolve()
    output = args.output.expanduser().resolve() if args.output else source.with_name(source.stem.replace("导入数据", "合并校对稿") + ".docx")
    count = build(source, output)
    print(f"已生成：{output}（{count} 条故事）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
